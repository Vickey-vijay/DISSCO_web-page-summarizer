"""Builds a Word document from Markdown chapter files, in the BITS report style.

Features: cover, certificate, progress page (optional), automatic table of contents, list of
figures, list of tables, numbered chapters / sections, numbered captions (real Word SEQ fields),
tables, code blocks, bullet and numbered lists, call-out quotes, page numbers.

Markdown subset understood in chapter files:
    # Chapter title                 -> Heading 1 (new page), auto-numbered "1", "2", ...
    ## Section / ### Subsection     -> Heading 2 / 3, auto-numbered "1.1", "1.1.1"
    # [nonum] References            -> Heading 1 without a number (references, appendices)
    plain paragraphs, **bold**, *italic*, `code`
    - bullet / * bullet             -> bullet list
    1. numbered item                -> numbered list (numbers are written as text)
    > quote                         -> shaded call-out paragraph
    ```text ... ```                -> code block
    Table: caption text             -> caption for the pipe table that follows
    | a | b |  / |---|---|         -> table
    ![Caption text](path/to/img.png) -> centred figure with numbered caption

Usage (see build_report.py): build_document(config_dict, out_path)
After building, run word_tools.refresh_and_export() so Word fills in the table of contents,
the lists of figures / tables and the page numbers.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

FONT = "Times New Roman"
MONO = "Consolas"
BLUE = RGBColor(0x1F, 0x3A, 0x8A)


# --------------------------------------------------------------------------- low-level helpers
def _set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _shade_paragraph(paragraph, hex_fill: str, left_border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    p_pr.append(shd)
    if left_border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), left_border)
        borders.append(left)
        p_pr.append(borders)


def _add_field(paragraph, instruction: str, cached_text: str = "") -> None:
    """Insert a complex Word field (TOC, SEQ, PAGE ...) with cached display text."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)

    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run2._r.append(instr)

    run3 = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run3._r.append(sep)

    run4 = paragraph.add_run(cached_text)

    run5 = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run5._r.append(end)
    return run4


def _set_font(style, name: str, size: float, bold=None, italic=None, color=None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(attr), name)
    # Built-in heading styles point at theme fonts (Calibri); drop those so the font above wins.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if r_fonts.get(qn(attr)) is not None:
            del r_fonts.attrib[qn(attr)]


# --------------------------------------------------------------------------- document setup
def _prepare_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.orientation = WD_ORIENT.PORTRAIT
    section.left_margin = section.right_margin = Cm(2.54)
    section.top_margin = section.bottom_margin = Cm(2.3)

    styles = doc.styles
    normal = styles["Normal"]
    _set_font(normal, FONT, 12)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (("Heading 1", 18, 0, 14), ("Heading 2", 14, 14, 6), ("Heading 3", 12.5, 10, 4)):
        st = styles[name]
        _set_font(st, FONT, size, bold=True, color=BLUE if name != "Heading 3" else RGBColor(0x1E, 0x29, 0x3B))
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True

    cap = styles["Caption"]
    _set_font(cap, FONT, 10.5, bold=True, italic=False, color=RGBColor(0x33, 0x41, 0x55))
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in ("List Bullet", "List Number"):
        _set_font(styles[name], FONT, 12)

    if "Code Block" not in [s.name for s in styles]:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        _set_font(code, MONO, 9)
        code.paragraph_format.line_spacing = 1.0
        code.paragraph_format.space_after = Pt(0)
        code.paragraph_format.left_indent = Cm(0.3)

    for name in ("TOC 1", "TOC 2", "TOC 3", "Table of Figures"):
        try:
            _set_font(styles[name], FONT, 11.5 if name != "TOC 3" else 11)
        except KeyError:
            pass


def _footer_and_header(doc: Document, short_title: str) -> None:
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True  # cover page stays clean
    header_p = sec.header.paragraphs[0]
    header_p.text = short_title
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in header_p.runs:
        r.font.size = Pt(9)
        r.font.name = FONT
        r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    footer_p = sec.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = _add_field(footer_p, "PAGE", "1")
    run.font.size = Pt(10)


# --------------------------------------------------------------------------- inline markdown
_INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def _add_inline(paragraph, text: str, size: float | None = None, mono_size: float = 10.5) -> None:
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), MONO)
            run.font.size = Pt(mono_size)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


# --------------------------------------------------------------------------- the builder
class Builder:
    def __init__(self, doc: Document, base_dir: Path):
        self.doc = doc
        self.base_dir = base_dir
        self.h = [0, 0, 0]
        self.figure_no = 0
        self.table_no = 0
        self.pending_table_caption: str | None = None
        self.warnings: list[str] = []

    # ---- headings
    def heading(self, level: int, text: str) -> None:
        numbered = True
        m = re.match(r"\[nonum\]\s*(.*)", text)
        if m:
            numbered, text = False, m.group(1)
        if numbered:
            self.h[level - 1] += 1
            for i in range(level, 3):
                self.h[i] = 0
            number = ".".join(str(n) for n in self.h[:level])
            text = f"{number}  {text}" if level > 1 else f"{number}.  {text}"
        self.doc.add_heading(text, level=level)

    # ---- figures and tables
    def figure(self, caption: str, rel_path: str) -> None:
        path = (self.base_dir / rel_path).resolve()
        if not path.exists():
            self.warnings.append(f"missing image: {rel_path}")
            p = self.doc.add_paragraph(f"[missing image: {rel_path}]")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return
        with Image.open(path) as im:
            w, h = im.size
        max_w, max_h = Cm(15.8), Cm(19.0)
        width = max_w
        if h / w * max_w > max_h:
            width = int(max_h * w / h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(str(path), width=width)
        self.figure_no += 1
        cap = self.doc.add_paragraph(style="Caption")
        cap.add_run("Figure ")
        _add_field(cap, "SEQ Figure \\* ARABIC", str(self.figure_no))
        cap.add_run(f": {caption}")

    def table(self, rows: list[list[str]], caption: str | None) -> None:
        if caption:
            self.table_no += 1
            cap = self.doc.add_paragraph(style="Caption")
            cap.paragraph_format.keep_with_next = True
            cap.add_run("Table ")
            _add_field(cap, "SEQ Table \\* ARABIC", str(self.table_no))
            cap.add_run(f": {caption}")
        ncols = max(len(r) for r in rows)
        table = self.doc.add_table(rows=len(rows), cols=ncols)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci in range(ncols):
                cell = table.cell(ri, ci)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                _add_inline(p, row[ci] if ci < len(row) else "", size=10, mono_size=9.5)
                if ri == 0:
                    for r in p.runs:
                        r.bold = True
                    _set_cell_shading(cell, "DBEAFE")
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---- markdown body
    def add_markdown(self, text: str) -> None:
        lines = text.replace("\r\n", "\n").split("\n")
        i = 0
        para: list[str] = []

        def flush_para():
            if para:
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                _add_inline(p, " ".join(s.strip() for s in para))
                para.clear()

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()

            if not stripped:
                flush_para()
                i += 1
                continue

            if stripped.startswith("```"):
                flush_para()
                i += 1
                block = []
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    block.append(lines[i].rstrip("\n"))
                    i += 1
                i += 1
                for k, code_line in enumerate(block or [""]):
                    p = self.doc.add_paragraph(style="Code Block")
                    _shade_paragraph(p, "F1F5F9")
                    run = p.add_run(code_line if code_line else " ")
                    run.font.name = MONO
                    run.font.size = Pt(9)
                self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
                continue

            m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
            if m:
                flush_para()
                self.heading(len(m.group(1)), m.group(2).strip())
                i += 1
                continue

            m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", stripped)
            if m:
                flush_para()
                self.figure(m.group(1).strip(), m.group(2).strip())
                i += 1
                continue

            m = re.match(r"^Table:\s*(.+)$", stripped)
            if m and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
                flush_para()
                self.pending_table_caption = m.group(1).strip()
                i += 1
                continue

            if stripped.startswith("|"):
                flush_para()
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                if rows:
                    self.table(rows, self.pending_table_caption)
                self.pending_table_caption = None
                continue

            if stripped.startswith(">"):
                flush_para()
                quote = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    quote.append(lines[i].strip().lstrip(">").strip())
                    i += 1
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.6)
                _shade_paragraph(p, "EFF6FF", left_border="2563EB")
                _add_inline(p, " ".join(quote))
                continue

            m = re.match(r"^[-*]\s+(.*)$", stripped)
            if m:
                flush_para()
                p = self.doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                _add_inline(p, m.group(1))
                i += 1
                continue

            m = re.match(r"^(\d+)[.)]\s+(.*)$", stripped)
            if m:
                flush_para()
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(1.0)
                p.paragraph_format.first_line_indent = Cm(-0.7)
                p.paragraph_format.space_after = Pt(3)
                _add_inline(p, f"{m.group(1)}.  {m.group(2)}")
                i += 1
                continue

            para.append(stripped)
            i += 1
        flush_para()


# --------------------------------------------------------------------------- front matter pages
def _centered(doc, text, size, bold=False, space_after=6, color=None, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _front_cover(doc: Document, cfg: dict) -> None:
    for _ in range(3):
        _centered(doc, "", 12)
    _centered(doc, cfg["title"], 26, bold=True, color=BLUE, space_after=10)
    _centered(doc, cfg.get("subtitle", ""), 15, italic=True, space_after=30)
    _centered(doc, cfg.get("report_kind", ""), 15, bold=True, space_after=40)
    for label, value in cfg.get("cover_rows", []):
        p = _centered(doc, "", 13, space_after=4)
        run = p.add_run(f"{label}: ")
        run.bold = True
        run.font.size = Pt(13)
        r2 = p.add_run(value)
        r2.font.size = Pt(13)
    for _ in range(3):
        _centered(doc, "", 12)
    _centered(doc, "BIRLA INSTITUTE OF TECHNOLOGY & SCIENCE, PILANI", 14, bold=True, space_after=4)
    _centered(doc, cfg.get("date_line", "[Insert Month and Year]"), 12)


def _certificate(doc: Document, cfg: dict) -> None:
    _page_break(doc)
    _centered(doc, "CERTIFICATE", 18, bold=True, space_after=18, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.6
    p.add_run(
        f"This is to certify that the project entitled \"{cfg['title']}\" submitted by "
        f"{cfg.get('student_name', '[Insert Student Name]')}, having BITS ID No. {cfg.get('bits_id', '[Insert BITS ID]')}, "
        f"for the partial fulfilment of the requirements of the {cfg.get('degree', '[Insert Degree Program]')} degree "
        "programme of BITS, embodies the bona fide work done by the student under my supervision."
    )
    for _ in range(4):
        doc.add_paragraph()
    for line in ("Signature of the Mentor", "Name, Designation, Organization & Location: [Insert]", "Place: [Insert]", "Date: [Insert]"):
        doc.add_paragraph(line)


def _progress_page(doc: Document, cfg: dict) -> None:
    prog = cfg.get("progress")
    if not prog:
        return
    _page_break(doc)
    _centered(doc, "Mid-Semester Progress Report", 18, bold=True, space_after=12, color=BLUE)
    for heading, items in (("Work Completed So Far", prog["completed"]), ("Remaining Tasks", prog["remaining"])):
        p = doc.add_paragraph()
        r = p.add_run(heading + ":")
        r.bold = True
        r.font.size = Pt(13)
        for group, bullets in items:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            p.add_run(group).bold = True
            for b in bullets:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)
                _add_inline(bp, b, size=11)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.add_run("Evaluation:").bold = True
    rows = [["EC No.", "Component", "Excellent", "Good", "Fair", "Poor"]] + [
        [str(i + 1), name, "", "", "", ""] for i, name in enumerate(["Design Project Outline", "Work Progress & Achievements", "Documentation & Presentation"])
    ]
    table = doc.add_table(rows=len(rows), cols=6)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = val
            for pr in cell.paragraphs:
                for rr in pr.runs:
                    rr.font.size = Pt(10.5)
                    rr.bold = ri == 0
            if ri == 0:
                _set_cell_shading(cell, "DBEAFE")
    doc.add_paragraph()
    doc.add_paragraph("Signature of the Mentor: ____________________        Date: ____________")


def _toc_pages(doc: Document) -> None:
    _page_break(doc)
    _centered(doc, "Table of Contents", 18, bold=True, space_after=12, color=BLUE)
    p = doc.add_paragraph()
    _add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
    _page_break(doc)
    _centered(doc, "List of Figures", 18, bold=True, space_after=12, color=BLUE)
    p = doc.add_paragraph()
    _add_field(p, 'TOC \\h \\z \\c "Figure"', "Right-click and choose Update Field to build the list of figures.")
    _page_break(doc)
    _centered(doc, "List of Tables", 18, bold=True, space_after=12, color=BLUE)
    p = doc.add_paragraph()
    _add_field(p, 'TOC \\h \\z \\c "Table"', "Right-click and choose Update Field to build the list of tables.")


# --------------------------------------------------------------------------- public entry point
def build_document(cfg: dict, out_path: str | Path) -> dict:
    """cfg keys: title, subtitle, report_kind, cover_rows [(label, value)], student_name, bits_id,
    degree, date_line, progress {completed, remaining}, short_title, base_dir, chapters [md paths],
    front_matter (bool: certificate + progress), toc (bool)."""
    base_dir = Path(cfg["base_dir"])
    doc = Document()
    _prepare_styles(doc)
    _footer_and_header(doc, cfg.get("short_title", cfg["title"]))
    if cfg.get("cover", True):
        _front_cover(doc, cfg)
    if cfg.get("front_matter", True):
        _certificate(doc, cfg)
        _progress_page(doc, cfg)
    if cfg.get("toc", True):
        _toc_pages(doc)

    builder = Builder(doc, base_dir)
    for chapter in cfg["chapters"]:
        text = (base_dir / chapter).read_text(encoding="utf-8-sig")
        builder.add_markdown(text)

    doc.core_properties.title = cfg["title"]
    doc.core_properties.author = cfg.get("student_name", "")
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return {"figures": builder.figure_no, "tables": builder.table_no, "warnings": builder.warnings}
